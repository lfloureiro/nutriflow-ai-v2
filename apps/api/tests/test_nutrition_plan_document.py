from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from app.services.nutrition_plan_document import (
    NutritionPlanDocumentError,
    extract_nutrition_plan_document,
)


def test_extract_utf8_text_document() -> None:
    result = extract_nutrition_plan_document(
        filename="plano.txt",
        content_type="text/plain",
        data="Pequeno-almoço: proteína mínimo 30 g.\nEvitar bebidas açucaradas.".encode(),
    )

    assert result.document_type == "text"
    assert "proteína mínimo 30 g" in result.source_text
    assert result.character_count == len(result.source_text)
    assert result.warnings == []


def test_extract_docx_includes_paragraphs_and_table_cells() -> None:
    document = Document()
    document.add_paragraph("Plano alimentar")
    document.add_paragraph("Ao pequeno-almoço, pelo menos 30 g de proteína.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Refeição"
    table.cell(0, 1).text = "Objetivo"
    table.cell(1, 0).text = "Almoço"
    table.cell(1, 1).text = "50 g proteína"
    buffer = BytesIO()
    document.save(buffer)

    result = extract_nutrition_plan_document(
        filename="plano.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=buffer.getvalue(),
    )

    assert result.document_type == "docx"
    assert "Plano alimentar" in result.source_text
    assert "30 g de proteína" in result.source_text
    assert "Refeição\tObjetivo" in result.source_text
    assert "Almoço\t50 g proteína" in result.source_text


def test_image_only_pdf_fails_closed_instead_of_returning_empty_text() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buffer = BytesIO()
    writer.write(buffer)

    with pytest.raises(NutritionPlanDocumentError, match="No readable text"):
        extract_nutrition_plan_document(
            filename="scan.pdf",
            content_type="application/pdf",
            data=buffer.getvalue(),
        )


def test_unsupported_document_type_is_rejected() -> None:
    with pytest.raises(NutritionPlanDocumentError, match="Unsupported document type"):
        extract_nutrition_plan_document(
            filename="plano.jpg",
            content_type="image/jpeg",
            data=b"not-an-image",
        )
