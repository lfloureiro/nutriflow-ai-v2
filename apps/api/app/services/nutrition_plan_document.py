from io import BytesIO
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.schemas.nutrition_plan_document import NutritionPlanDocumentExtractionRead

EXTRACTOR_NAME = "nutriflow-document-text"
EXTRACTOR_VERSION = "v1"
MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
MAX_SOURCE_CHARACTERS = 100_000
MAX_PDF_PAGES = 60
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}


class NutritionPlanDocumentError(ValueError):
    pass


def _clean_text(value: str) -> str:
    lines = [line.rstrip() for line in value.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    compact: list[str] = []
    previous_blank = False
    for line in lines:
        is_blank = not line.strip()
        if is_blank and previous_blank:
            continue
        compact.append(line)
        previous_blank = is_blank
    return "\n".join(compact).strip()


def _validate_text(text: str) -> str:
    cleaned = _clean_text(text)
    if not cleaned:
        raise NutritionPlanDocumentError(
            "No readable text was found in the document. Scanned/image-only files require image extraction."
        )
    if len(cleaned) > MAX_SOURCE_CHARACTERS:
        raise NutritionPlanDocumentError(
            f"Extracted document text exceeds the {MAX_SOURCE_CHARACTERS} character limit."
        )
    return cleaned


def _extract_text_file(data: bytes) -> tuple[str, list[str]]:
    try:
        return data.decode("utf-8-sig"), []
    except UnicodeDecodeError as exc:
        raise NutritionPlanDocumentError("Text documents must use UTF-8 encoding.") from exc


def _extract_docx(data: bytes) -> tuple[str, list[str]]:
    try:
        document = Document(BytesIO(data))
    except (PackageNotFoundError, ValueError) as exc:
        raise NutritionPlanDocumentError("The Word document could not be read as DOCX.") from exc

    blocks: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            blocks.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append("\t".join(cells))
    return "\n".join(blocks), []


def _extract_pdf(data: bytes) -> tuple[str, list[str]]:
    try:
        reader = PdfReader(BytesIO(data), strict=False)
    except (PdfReadError, ValueError, OSError) as exc:
        raise NutritionPlanDocumentError("The PDF document could not be read.") from exc

    if reader.is_encrypted:
        raise NutritionPlanDocumentError("Encrypted/password-protected PDFs are not supported.")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise NutritionPlanDocumentError(
            f"PDF documents are limited to {MAX_PDF_PAGES} pages for plan import."
        )

    warnings: list[str] = []
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except (PdfReadError, ValueError) as exc:
            raise NutritionPlanDocumentError(
                f"Text extraction failed on PDF page {index}."
            ) from exc
        if text.strip():
            pages.append(text)
        else:
            warnings.append(f"pdf_page_without_text:{index}")
    return "\n\n".join(pages), warnings


def extract_nutrition_plan_document(
    *,
    filename: str,
    content_type: str | None,
    data: bytes,
) -> NutritionPlanDocumentExtractionRead:
    safe_name = Path(filename or "document").name
    suffix = Path(safe_name).suffix.casefold()
    if suffix not in SUPPORTED_SUFFIXES:
        allowed = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise NutritionPlanDocumentError(f"Unsupported document type. Allowed: {allowed}.")
    if not data:
        raise NutritionPlanDocumentError("The uploaded document is empty.")
    if len(data) > MAX_DOCUMENT_BYTES:
        raise NutritionPlanDocumentError("The uploaded document exceeds the 10 MB limit.")

    if suffix == ".pdf":
        raw_text, warnings = _extract_pdf(data)
        document_type = "pdf"
    elif suffix == ".docx":
        raw_text, warnings = _extract_docx(data)
        document_type = "docx"
    else:
        raw_text, warnings = _extract_text_file(data)
        document_type = "text"

    source_text = _validate_text(raw_text)
    return NutritionPlanDocumentExtractionRead(
        filename=safe_name,
        content_type=content_type,
        document_type=document_type,
        source_text=source_text,
        character_count=len(source_text),
        extractor_name=EXTRACTOR_NAME,
        extractor_version=EXTRACTOR_VERSION,
        warnings=warnings,
    )
